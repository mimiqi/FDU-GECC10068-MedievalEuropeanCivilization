from pathlib import Path
from copy import deepcopy
import hashlib, json, re
from docx import Document
from docx.oxml.ns import qn

ROOT=Path(r'D:\text\02_中世纪欧洲文明\讨论课')
WORK=ROOT/'.制作过程'/'第八页与讲稿同步'
SOURCE=ROOT/'交付文件'/'查理曼_12分钟汇报讲稿.docx'
EXPECTED='fe4f9a7c6f82d9f726d65b32aa77e95d6ce8a186509b3d7049e61f578b2e2c22'
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==EXPECTED, 'Source changed'
data=json.loads((WORK/'content.json').read_text(encoding='utf-8'))
doc=Document(SOURCE)
paras=doc.paragraphs

def assign(p,text):
    props=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    starts=[deepcopy(e) for e in p._p.findall(qn('w:commentRangeStart'))]
    ends=[deepcopy(e) for e in p._p.findall(qn('w:commentRangeEnd'))]
    refs=[deepcopy(r) for r in p._p.findall(qn('w:r')) if r.find(qn('w:commentReference')) is not None]
    p.clear()
    for el in starts:p._p.append(el)
    run=p.add_run(text)
    if props is not None:run._r.insert(0,props)
    for el in ends+refs:p._p.append(el)

blocks={
 1:(9,10,[11],12),2:(13,14,[15],16),3:(17,18,[19,20],21),
 4:(24,25,[26,27],28),5:(29,30,[31,32,33],34),6:(37,38,[39,40,41],42),
 7:(43,44,[45,46],47),8:(50,51,[52,53],54),9:(55,56,[57,58],59),10:(62,63,[64,65],66)
}
short_refs={
 2:'Nelson, King and Emperor, PDF 251、405、407—408页；巴尔贝罗《查理大帝：欧洲之父》210、238—239页（PDF 219、247—248）。',
 6:'布洛赫《封建社会》上卷266—269页（PDF 288—291）；巴尔贝罗《查理大帝：欧洲之父》146—149页（PDF 155—158）。',
 7:'巴尔贝罗《查理大帝：欧洲之父》231—237页（PDF 240—246）：办学232—233页，校订与传播231、234页，小写字体236—237页。',
 8:'巴尔贝罗《查理大帝：欧洲之父》113、234页；艾因哈德《查理大帝传》28页；布洛赫《封建社会》上卷266—269页；Nelson, King and Emperor, PDF 408、515页。'
}
durations=[20,70,40,70,110,110,70,80,70,50]
for s,duration in zip(data['slides'],durations):
    h,t,body,ref=blocks[s['oldNumber']]
    assert len(body)==len(s['script'])
    assign(paras[h],f"第{s['n']:02d}页  {s['title']}")
    assign(paras[t],f"建议用时：{duration}秒  |  累计时间：{s['time']}")
    for idx,text in zip(body,s['script']):assign(paras[idx],text)
    reftext=short_refs.get(s['n'],s['ref'])
    reftext=re.sub(r'^(立论依据：|综合：|结论依据：)','',reftext)
    assign(paras[ref],'参考：'+reftext)

# Follow the user's actual slide order: Roman traditions precede the analytical framework.
anchor=paras[21]._p
for p in paras[13:17]:
    anchor.addnext(p._p)
    anchor=p._p

han=len(re.findall(r'[\u4e00-\u9fff]',''.join(t for s in data['slides'] for t in s['script'])))
assign(paras[3],'配套PPT：查理曼_政治治理视角_12分钟汇报_同步修订版.pptx\n小班：[待填写]    小组：第八小组    主讲：熊礼学\nPPT：张祺昀\n讨论稿：张元赫、张芯源、朱硕、张智语')
assign(paras[5],f'主讲为第1—10页，安排11分30秒内容，另留30秒用于停顿、换页或现场调整。第11页为书目，不计入主讲。经济内容保留在本讲稿的答辩备用部分。正文约{round(han/100)*100}个汉字，请按主讲者语速排练。')
assign(paras[8],'汇报正文  罗马传统与讨论思路')
assign(paras[49],'汇报正文  地区联系与君主性质')
assign(paras[67],'结束提示：主讲到此结束。答辩时可保持第10页结论画面；经济问题参阅下方备用发言。')
assign(paras[69],'经济答辩备用（约50秒）')
assign(paras[73],'我们主要根据实际治理来判断。他以法兰克地区为统治重心，通过巡察使监督伯爵，借领主组织人员，又依靠教会的读写人才和跨地区往来。这些做法显示了他对当时统治条件的回应。地理位置只是比较的一部分；只有把它与治理方式结合，才能说明我们为什么更倾向于称他为新型欧洲君主。')
assign(paras[74],'依据：巴尔贝罗《查理大帝：欧洲之父》113、206—210、231—237页；布洛赫《封建社会》上卷266—269页；艾因哈德《查理大帝传》28页。')
assign(paras[87],'如果把“欧洲之父”理解成现代欧洲的创造者，就会夸大影响。我们说的“欧洲性”，是西欧、中欧部分地区与意大利北部之间的政治联系，以及拉丁基督教文化中的人员和书籍往来。阿尔昆来自不列颠，也说明文化联系可以超出帝国疆界，不能用学者的活动范围推定查理曼的领土。')
assign(paras[88],'依据：巴尔贝罗《查理大帝：欧洲之父》113、234页；艾因哈德《查理大帝传》28页；Nelson, King and Emperor, PDF 408、515页。')
assign(paras[96],'本稿结合本组四份讨论稿及2022年两份参考材料。张元赫稿中的建筑、地理重心和继承仪式例子，核对了艾因哈德28、31页、巴尔贝罗113页与Nelson PDF 497—498页；张智语稿中的文化改革和帝国分裂后的影响，核对了巴尔贝罗231—237页及皮朗290页。第8页另以艾因哈德28页的宫廷学者为例，说明跨地区文化往来。艾因哈德是宫廷传记作者，其叙述还需结合后世研究理解。')
assign(paras[105],paras[105].text.replace('231—237','231—239'))
assign(paras[111],paras[111].text.replace('204、','204、251、'))

for row,s,duration in zip(doc.tables[0].rows[1:],data['slides'],durations):
    assign(row.cells[0].paragraphs[0],str(s['n']))
    assign(row.cells[1].paragraphs[0],s['title'].replace('\n',''))
    assign(row.cells[2].paragraphs[0],f'{duration}秒')

comment_updates={
 0:'罗马传统包括帝国文书称号、君主召集教会会议及具体案件采用罗马法。依据Nelson PDF 251、405、407—408页和巴尔贝罗210、238—239页；不推论全帝国统一采用罗马法。',
 1:'第8页由地理比较推进到地方治理和文化往来。地理据巴尔贝罗113页，学者来源据艾因哈德28页，书籍流通据巴尔贝罗234页，亚琛及非目的论限定据Nelson PDF 408、515页；文化联系不等于政治疆界。',
 4:'教会办学、阿尔昆校订与抄本传播、加洛林小写字体分别依据巴尔贝罗232—233、231与234、236—237页。保留训令与实际落实的区别，不把字体发明归于查理曼个人。'
}
for c in doc.comments:
    if c.comment_id in comment_updates:
        c.paragraphs[0].text=comment_updates[c.comment_id]
        for p in c.paragraphs[1:]:p._p.getparent().remove(p._p)

output=WORK/'讲稿修订候选.docx'
doc.save(output)
# All main-speech paragraphs must match the PowerPoint notes, in physical slide order.
check=Document(output)
texts=[p.text for p in check.paragraphs]
last=-1
for s in data['slides']:
    heading=f"第{s['n']:02d}页  {s['title']}"
    idx=texts.index(heading)
    assert idx>last
    assert texts[idx+2:idx+2+len(s['script'])]==s['script']
    last=idx
assert '第12页' not in '\n'.join(texts)
comment_ids={str(c.comment_id) for c in check.comments}
for tag in ['commentRangeStart','commentRangeEnd','commentReference']:
    ids=[e.get(qn('w:id')) for e in check.element.iter(qn('w:'+tag))]
    assert set(ids)==comment_ids and len(ids)==len(comment_ids),(tag,ids,comment_ids)
(WORK/'manuscript-validation.json').write_text(json.dumps({'han':han,'mainSpeechMatchesPpt':True,'comments':len(comment_ids),'sourceSha256':EXPECTED,'candidateSha256':hashlib.sha256(output.read_bytes()).hexdigest()},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'han':han,'comments':len(comment_ids),'mainSpeechMatchesPpt':True}))
