from pathlib import Path
from copy import deepcopy
import json
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT=Path(r'D:\text\02_中世纪欧洲文明\讨论课')
WORK=ROOT/'.制作过程'/'修订'
old=json.loads((ROOT/'.制作过程'/'content.json').read_text(encoding='utf-8'))
new=json.loads((WORK/'content.revised.json').read_text(encoding='utf-8'))
doc=Document(ROOT/'交付文件'/'查理曼_12分钟汇报讲稿.docx')

def assign(p,text):
    props=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear();r=p.add_run(text)
    if props is not None:r._r.insert(0,props)

def reftext(text):
    for lead in ['立论依据：','综合：','结论依据：']:
        if text.startswith(lead):text=text[len(lead):]
    return text

mapping={old['subtitle']:new['subtitle'],
 '“新型”指统治资源的重组，“欧洲性”限定为特定地区之间的政治与拉丁基督教联系。文中引述均按资料概括转述，未将二手研究写成未经核验的史料原文。':'本稿从地方司法、附庸关系和教会教育入手，讨论查理曼如何治理帝国。引用以转述为主，相关书名和页码列在各段之后，汇报时不用朗读。',
 '汇报正文 · 提出问题与承认延续':'汇报正文  问题与罗马传统',
 '汇报正文 · 从征服走向治理':'汇报正文  领土扩张与地方治理',
 '汇报正文 · 动员与沟通的资源':'汇报正文  附庸关系与教会教育',
 '汇报正文 · “欧洲性”与最终判断':'汇报正文  帝国的地理重心与君主性质',
 '汇报正文 · 结论':'汇报正文  结论',
 '答辩准备 · 证据与解释的边界':'答辩准备  制度与史料',
 '艾因哈德为宫廷传记作者；巴尔贝罗、布洛赫与Nelson属于后世研究。注意区分当时叙述与后世解释。本稿结合本组两份讨论稿及2022年两份参考材料的思路，具体史实与页码以课程所供参考书核对。':'本稿结合本组三份讨论稿及2022年两份参考材料。新增的建筑、地理重心和继承仪式例子，分别核对了艾因哈德28、31页，巴尔贝罗113页与Nelson PDF 497—498页。艾因哈德是宫廷传记作者，其叙述还需结合后世研究理解。'
}
for a,b in zip(old['slides'],new['slides']):
    if a['n']<=10:
        mapping[f"第{a['n']:02d}页  {a['title'].replace(chr(10),'')}"]=f"第{b['n']:02d}页  {b['title'].replace(chr(10),'')}"
        mapping['本页依据：'+a['ref']]='参考：'+reftext(b['ref'])
    mapping[a['ref']]=b['ref']
    assert len(a['script'])==len(b['script'])
    for prev,now in zip(a['script'],b['script']):mapping[prev]=now
for a,b in zip(old['qa'],new['qa']):
    mapping[a['q']]=b['q'];mapping[a['a']]=b['a'];mapping['依据：'+a['ref']]='依据：'+b['ref']
for a,b in zip(old['bibliography'],new['bibliography']):mapping[a['pages']]=b['pages']
count=0
for p in doc.paragraphs:
    if p.text in mapping:assign(p,mapping[p.text]);count+=1
    elif '正文约2500个汉字' in p.text:assign(p,p.text.replace('正文约2500个汉字','正文约2400个汉字'))
    elif p.text.startswith('建议时段 '):assign(p,p.text.split('；')[0].replace('  ｜  ','，')+'。')
for row,s in zip(doc.tables[0].rows[1:],new['slides'][:10]):
    assign(row.cells[1].paragraphs[0],s['title'].replace('\n',''))

# Put the final two Q&A topics together so they do not crowd the previous page.
europe=next(p for p in doc.paragraphs if p.text==new['qa'][4]['q'])
europe.paragraph_format.page_break_before=True
material=next(p for p in doc.paragraphs if p.text=='材料使用说明')
def insert_before(anchor,text,style,model=None):
    el=OxmlElement('w:p');anchor._p.addprevious(el);p=Paragraph(el,anchor._parent)
    if model is not None and model._p.pPr is not None:el.insert(0,deepcopy(model._p.pPr))
    p.style=style
    if model is not None and model.runs:
        run=p.add_run(text)
        if model.runs[0]._r.rPr is not None:run._r.insert(0,deepcopy(model.runs[0]._r.rPr))
    else:p.add_run(text)
    p.paragraph_format.page_break_before=False
    return p
body_model=next(p for p in doc.paragraphs if p.text==new['qa'][3]['a'])
ref_model=next(p for p in doc.paragraphs if p.text=='依据：'+new['qa'][3]['ref'])
extra=new['qa'][5]
insert_before(material,extra['q'],'Heading 2')
added=insert_before(material,extra['a'],'Normal',body_model)
insert_before(material,'依据：'+extra['ref'],'Normal',ref_model)

# Notes explain the three source additions without intruding on the speech text.
for p in doc.paragraphs:
    if p.text==new['slides'][2]['script'][0]:
        doc.add_comment(p.runs[0],'据张元赫讨论稿补入亚琛礼拜堂的例子，已核对艾因哈德《查理大帝传》28页（PDF 37）。',author='修订说明',initials='修')
    elif p.text==new['slides'][7]['script'][0]:
        doc.add_comment(p.runs[0],'据张元赫讨论稿补充地理重心的对照，已核对巴尔贝罗《查理大帝：欧洲之父》113页（PDF 122）。',author='修订说明',initials='修')
doc.add_comment(added.runs[0],'新增答辩例子。艾因哈德31页记述查理曼安排路易共同执政并取得皇帝称号；813年的日期参照Nelson PDF 497—498页。',author='修订说明',initials='修')

# Remove the Word template's decorative title rule while keeping all page geometry.
title_style=doc.styles['Title']
for node in list(title_style.element.findall('.//'+qn('w:pBdr'))):node.getparent().remove(node)
for p in doc.paragraphs:
    if p.style.name=='Title':
        for node in list(p._p.findall('.//'+qn('w:pBdr'))):node.getparent().remove(node)

doc.core_properties.subject='查理曼的地方治理、附庸关系、教会教育与罗马传统'
doc.save(WORK/'讲稿修订候选.docx')
print('Edited existing document:',count,'paragraph replacements, 3 source comments, 1 added Q&A.')
