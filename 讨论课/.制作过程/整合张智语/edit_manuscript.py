from pathlib import Path
from copy import deepcopy
import json, re
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement

ROOT = Path(r'D:\text\02_中世纪欧洲文明\讨论课')
WORK = ROOT / '.制作过程' / '整合张智语'
old = json.loads((ROOT / '.制作过程' / '修订' / 'content.revised.json').read_text(encoding='utf-8'))
new = deepcopy(old)
s7 = new['slides'][6]
s7['title'] = '教育与书写改革怎样帮助治理'
s7['ref'] = '巴尔贝罗《查理大帝：欧洲之父》231—237页（PDF 240—246）。'
s7['script'] = [
    '文化方面，查理曼支持的教育和抄书活动，是加洛林文艺复兴的重要内容。789年的训令要求教会办学，教授语法等知识。他也支持校订《圣经》，减少抄本中的错误。同时，加洛林小写字体逐渐传播，字形更清晰，便于辨认。这种字体由抄写员和学者逐渐发展起来，不能说是查理曼本人发明的。',
    '这些改革与政治有什么关系？处理文书、传达命令和解释教义，都需要熟悉拉丁文的人。教会学校培养这样的人员，较为一致的书写方式也方便不同地区阅读和使用文本。因此，文化建设除了服务信仰，也有助于帝国的日常治理。这里的变化，是查理曼进一步支持了教会教育和书写规范。'
]
s10 = new['slides'][9]
s10['ref'] = '结论依据：艾因哈德《查理大帝传》10—12页；巴尔贝罗《查理大帝：欧洲之父》113、206—210、231—237页；布洛赫《封建社会》上卷266—269页；皮朗《穆罕默德和查理曼》290页（PDF 322）。'
s10['script'] = [
    '回到讨论题，查理曼的主要功绩，是扩大法兰克统治，整顿地方司法，并支持教会教育和拉丁文化。帝国后来发生分裂，但教会之间的往来、书籍抄写和共同的拉丁文化并未因此消失。评价他的影响，不能只看政治统一维持了多久。',
    '因此，我们更倾向于称他为新型欧洲君主。他继承了罗马皇帝的理想，统治的重心在法兰克地区，也需要地方贵族和教会合作。前面讨论的司法、附庸关系和文化教育，正好说明了这种统治的特点。谢谢大家。'
]
new['bibliography'][1]['pages'] = '书内113、146—149、206—210、229、231—237、271—272页；PDF页码＝书内页码＋9；地图PDF 5。'
new['bibliography'][4]['details'] = '[比] 亨利·皮朗著。用于讨论政治分裂后的文化延续，以及经济答辩的史学争议背景。'
new['bibliography'][4]['pages'] = '书内290页（PDF 322）；译者前言4页（PDF 10）。'
new['qa'].append({
    'q': '帝国后来分裂，是否说明查理曼的建设失败了？',
    'a': '要分别看政治统一和文化影响。分裂说明查理曼没有建立起能长期维持统一的政治秩序，但不能据此否定所有成果。皮朗在第290页强调，政治统一结束后，共同文化仍可延续。书籍的保存与流通，以及拉丁文教育和书写规范，就是理解这种影响的具体例子。不过，这不意味着巡察使等制度原样延续，也不能把后来的欧洲都归功于查理曼。',
    'ref': '皮朗《穆罕默德和查理曼》290页（PDF 322）；巴尔贝罗《查理大帝：欧洲之父》232—237页（PDF 241—246）。'
})
(WORK / 'content.revised.json').write_text(json.dumps(new, ensure_ascii=False, indent=2), encoding='utf-8')
doc = Document(ROOT / '交付文件' / '查理曼_12分钟汇报讲稿.docx')

def assign(p, text):
    props = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    run = p.add_run(text)
    if props is not None:
        run._r.insert(0, props)

def reftext(text):
    for lead in ['立论依据：', '综合：', '结论依据：']:
        if text.startswith(lead):
            return text[len(lead):]
    return text

mapping = {}
for idx in [6, 9]:
    a, b = old['slides'][idx], new['slides'][idx]
    mapping[f"第{a['n']:02d}页  {a['title']}"] = f"第{b['n']:02d}页  {b['title']}"
    mapping['参考：' + reftext(a['ref'])] = '参考：' + reftext(b['ref'])
    for prev, now in zip(a['script'], b['script']):
        mapping[prev] = now
for idx in [1, 4]:
    for key in ['pages', 'details']:
        mapping[old['bibliography'][idx][key]] = new['bibliography'][idx][key]
for p in doc.paragraphs:
    if p.text in mapping:
        assign(p, mapping[p.text])
    elif p.text.startswith('本稿结合本组三份讨论稿'):
        assign(p, '本稿结合本组四份讨论稿及2022年两份参考材料。张元赫稿中的建筑、地理重心和继承仪式例子，核对了艾因哈德28、31页、巴尔贝罗113页与Nelson PDF 497—498页；张智语稿中的文化改革和帝国分裂后的影响，核对了巴尔贝罗231—237页及皮朗290页。艾因哈德是宫廷传记作者，其叙述还需结合后世研究理解。')

for row, s in zip(doc.tables[0].rows[1:], new['slides'][:10]):
    if row.cells[1].text != s['title'].replace('\n', ''):
        assign(row.cells[1].paragraphs[0], s['title'].replace('\n', ''))

material = next(p for p in doc.paragraphs if p.text == '材料使用说明')
body_model = next(p for p in doc.paragraphs if p.text == new['qa'][3]['a'])
ref_model = next(p for p in doc.paragraphs if p.text == '依据：' + new['qa'][3]['ref'])

def insert_before(anchor, text, style, model=None):
    el = OxmlElement('w:p')
    anchor._p.addprevious(el)
    p = Paragraph(el, anchor._parent)
    if model is not None and model._p.pPr is not None:
        el.insert(0, deepcopy(model._p.pPr))
    p.style = style
    run = p.add_run(text)
    if model is not None and model.runs and model.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(model.runs[0]._r.rPr))
    p.paragraph_format.page_break_before = False
    return p

extra = new['qa'][-1]
insert_before(material, extra['q'], 'Heading 2')
added = insert_before(material, extra['a'], 'Normal', body_model)
insert_before(material, '依据：' + extra['ref'], 'Normal', ref_model)
doc.add_comment(added.runs[0], '吸收张智语稿对政治分裂与长期影响的区分，核对皮朗290页及巴尔贝罗232—237页。限定为文化联系的延续，没有推论所有政治制度持续存在。', author='修订说明', initials='修')
for p in doc.paragraphs:
    if p.text == s7['script'][0]:
        doc.add_comment(p.runs[0], '据张智语稿补充办学、文本校订与小写字体的具体例子，已核对巴尔贝罗231—237页。保留训令要求与实际落实的区别，也未将字体的发明归于查理曼个人。', author='修订说明', initials='修')

han = len(re.findall(r'[\u4e00-\u9fff]', ''.join(t for s in new['slides'][:10] for t in s['script'])))
for p in doc.paragraphs:
    if '正文约2400个汉字' in p.text:
        assign(p, p.text.replace('正文约2400个汉字', f'正文约{round(han/100)*100}个汉字'))
doc.save(WORK / '讲稿修订候选.docx')
print(json.dumps({'mainSpeechHanCharacters': han, 'slides': len(new['slides']), 'qa': len(new['qa'])}))
