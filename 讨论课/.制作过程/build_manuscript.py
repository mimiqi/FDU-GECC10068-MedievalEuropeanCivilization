from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'D:\text\02_中世纪欧洲文明\讨论课')
DATA=json.loads((ROOT/'.制作过程'/'content.json').read_text(encoding='utf-8'))
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.68)
sec.left_margin=Inches(.86);sec.right_margin=Inches(.86)
sec.header_distance=Inches(.3);sec.footer_distance=Inches(.3)

def setfont(style,size,bold=False,color='171717'):
    style.font.name='Microsoft YaHei'
    style.font.size=Pt(size)
    style.font.bold=bold
    style.font.color.rgb=RGBColor.from_string(color)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')

setfont(doc.styles['Normal'],12.5)
doc.styles['Normal'].paragraph_format.line_spacing=Pt(19)
doc.styles['Normal'].paragraph_format.space_after=Pt(9)
for name,size in [('Title',24),('Heading 1',17),('Heading 2',14)]:
    setfont(doc.styles[name],size,True)
    doc.styles[name].paragraph_format.space_before=Pt(9)
    doc.styles[name].paragraph_format.space_after=Pt(9)
    doc.styles[name].paragraph_format.keep_with_next=True
for name in ['Header','Footer']:
    setfont(doc.styles[name],9,False,'666666')
for style in doc.styles:
    if style.type == 1:
        pr=style.element.get_or_add_pPr()
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pr.append(snap)

header=sec.header.paragraphs[0];header.text='中世纪欧洲文明  ·  小组讨论汇报'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('查理曼：政治治理视角  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def para(text,style=None):
    return doc.add_paragraph(text,style)
def note(text):
    p=para(text);p.paragraph_format.space_after=Pt(10);p.paragraph_format.line_spacing=1.15
    for run in p.runs:
        run.font.size=Pt(9);run.font.color.rgb=RGBColor.from_string('666666')
    return p
def title(text):
    doc.add_heading(text,1)
def speech(n):
    s=DATA['slides'][n-1]
    p=doc.add_heading(f"第{n:02d}页  {s['title'].replace(chr(10),'')}",2)
    note(f"建议时段 {s['time']}  ｜  {s['seconds']}秒；以下正文朗读，书目提示不朗读。")
    for text in s['script']:
        p=para(text);p.paragraph_format.first_line_indent=Pt(25)
    note('本页依据：'+s['ref'])
def newpage():doc.add_page_break()
def shade(cell,color):
    pr=cell._tc.get_or_add_tcPr();el=OxmlElement('w:shd');el.set(qn('w:fill'),color);pr.append(el)

doc.add_heading(DATA['title'],0)
para(DATA['subtitle'])
para('12分钟汇报讲稿 · 含逐页提示与答辩准备')
note('配套PPT：查理曼_政治治理视角_12分钟汇报.pptx\n小班：[待填写]    小组：[待填写]    主讲：[待填写]\n成员及分工：[请与PPT首页同步填写]')
para('使用说明', 'Heading 2')
para('主讲为第1—10页，安排11分30秒内容，另留30秒用于停顿、换页或现场调整。第11页为书目，第12页为经济答辩备用，均不计入主讲。正文约2500个汉字；请根据主讲者语速排练。')
note('“新型”指统治资源的重组，“欧洲性”限定为特定地区之间的政治与拉丁基督教联系。文中引述均按资料概括转述，未将二手研究写成未经核验的史料原文。')
table=doc.add_table(rows=1,cols=3);table.style='Table Grid'
for c,text in zip(table.rows[0].cells,['页码','内容','建议时段']):
    c.text=text;shade(c,'EDEAE3')
for s in DATA['slides'][:10]:
    cells=table.add_row().cells
    for c,text in zip(cells,[str(s['n']),s['title'].replace('\n',''),s['time']]):c.text=text
for row in table.rows:
    row.cells[0].width=Inches(.5);row.cells[1].width=Inches(4.3);row.cells[2].width=Inches(1.5)
    for c in row.cells:
        for p in c.paragraphs:
            p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
            p.paragraph_format.line_spacing=1.1
            for r in p.runs:r.font.size=Pt(10)

newpage();title('汇报正文 · 提出问题与承认延续')
for n in [1,2,3]:speech(n)
newpage();title('汇报正文 · 从征服走向治理')
for n in [4,5]:speech(n)
newpage();title('汇报正文 · 动员与沟通的资源')
for n in [6,7]:speech(n)
newpage();title('汇报正文 · “欧洲性”与最终判断')
for n in [8,9]:speech(n)
newpage();title('汇报正文 · 结论')
speech(10)
note('结束提示：主讲到此结束。保持第10页结论画面，提问涉及经济时再切换到第12页。')

title('以下为答辩备用，不计入主讲')
doc.add_heading('第12页备用发言（约50秒）',2)
for text in DATA['slides'][11]['script']:para(text)
note(DATA['slides'][11]['ref'])
for item in DATA['qa'][:1]:
    doc.add_heading(item['q'],2);para(item['a']);note('依据：'+item['ref'])
newpage();title('答辩准备 · 证据与解释的边界')
for item in DATA['qa'][1:]:
    doc.add_heading(item['q'],2);para(item['a']);note('依据：'+item['ref'])
doc.add_heading('材料使用说明',2)
note('艾因哈德为宫廷传记作者；巴尔贝罗、布洛赫与Nelson属于后世研究。注意区分当时叙述与后世解释。本稿结合本组两份讨论稿及2022年两份参考材料的思路，具体史实与页码以课程所供参考书核对。')

newpage();title('参考书目与定位页码')
para('本稿以“讨论课”目录内课程所供PDF为准。中文书优先使用书内印刷页码；PDF页码从文件第一页起计数。英文Nelson一书统一使用PDF顺序页码。')
for i,b in enumerate(DATA['bibliography'],1):
    doc.add_heading(f"{i}. {b['short']} {b['title']}",2)
    para(b['details']);note(b['pages'])
doc.add_heading('页码复核示例',2)
para('司法改革：巴尔贝罗书内209页＝PDF 218页。\n附庸动员：布洛赫上卷267页＝合并PDF 289页。\n加冕叙述：艾因哈德书内30页＝PDF 39页。\n原书地图：巴尔贝罗PDF 5，无书内印刷页码。')

doc.core_properties.title=DATA['title']+'——12分钟汇报讲稿'
doc.core_properties.subject='政治治理、司法改革、附庸关系与拉丁基督教网络'
doc.core_properties.author='小组成员待填写'
doc.core_properties.keywords='查理曼,中世纪欧洲文明,12分钟汇报,政治治理'
out=ROOT/'交付文件'/'查理曼_12分钟汇报讲稿.docx'
out.parent.mkdir(exist_ok=True)
doc.save(out)
print(out)
print('Main speech Han characters:',sum(len(re.findall(r'[\u4e00-\u9fff]',''.join(s['script']))) for s in DATA['slides'][:10]))
